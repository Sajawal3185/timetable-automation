"""Tests for timetable_import.py: parsing helpers, flat and grid sheets, validation, loading."""
from datetime import time

import pytest
from openpyxl import Workbook
from sqlalchemy import select

from app.models.room import Room
from app.models.timetable_entry import TimetableEntry
from app.services import timetable_import as ti


# ---------------------------------------------------------------- small helpers
@pytest.mark.parametrize("text, expected", [
    ("8:30", time(8, 30)), ("08.30", time(8, 30)), ("8:30 AM", time(8, 30)),
    ("2 pm", time(14, 0)), ("2:15 PM", time(14, 15)), ("12:00", time(12, 0)),
    ("12 am", time(0, 0)), ("1:00", time(13, 0)),  # 1-7 with no am/pm = afternoon
    ("10:45", time(10, 45)),
])
def test_parse_time(text, expected):
    assert ti.parse_time(text) == expected


@pytest.mark.parametrize("text", ["", "abc", "25:00", "Monday"])
def test_parse_time_rejects_non_times(text):
    assert ti.parse_time(text) is None


@pytest.mark.parametrize("text, expected", [
    ("08:30 - 10:00", (time(8, 30), time(10, 0))),
    ("8:30-10:00", (time(8, 30), time(10, 0))),
    ("8:30 to 10:00", (time(8, 30), time(10, 0))),
    ("11:30 \u2013 1:00 PM", (time(11, 30), time(13, 0))),
    ("1:00-2:30", (time(13, 0), time(14, 30))),
])
def test_parse_time_range(text, expected):
    assert ti.parse_time_range(text) == expected


@pytest.mark.parametrize("text", ["", "8:30", "Course", "x - y"])
def test_parse_time_range_rejects(text):
    assert ti.parse_time_range(text) is None


@pytest.mark.parametrize("text, expected", [
    ("Mon", "Monday"), ("monday", "Monday"), ("TUES", "Tuesday"), ("Wed.", "Wednesday"),
    ("Thu", "Thursday"), ("thursday", "Thursday"), ("Fri", "Friday"), ("Sat", "Saturday"),
])
def test_normalize_day(text, expected):
    assert ti.normalize_day(text) == expected


@pytest.mark.parametrize("text", ["", "Lunch", "Mo", "Course"])
def test_normalize_day_rejects(text):
    assert ti.normalize_day(text) is None


def test_split_course():
    assert ti.split_course("CS-101 Programming Fundamentals") == ("CS-101", "Programming Fundamentals")
    assert ti.split_course("SE201 - Software Design") == ("SE201", "Software Design")
    assert ti.split_course("Islamic Studies") == ("", "Islamic Studies")


def test_parse_cell_three_lines():
    assert ti.parse_cell("CS101 Programming\nDr. Ali\nRoom 5") == ("CS101", "Programming", "Dr. Ali", "Room 5")


def test_parse_cell_pipe_separated():
    assert ti.parse_cell("Calculus | Ms. Sara | LT-2") == ("", "Calculus", "Ms. Sara", "LT-2")


def test_parse_cell_two_lines_detects_room_or_teacher():
    assert ti.parse_cell("Physics\nLab 3")[3] == "Lab 3"
    assert ti.parse_cell("Physics\nDr. Khan")[2] == "Dr. Khan"


@pytest.mark.parametrize("text", ["", "-", "Break", "LUNCH", "free"])
def test_parse_cell_empty(text):
    assert ti.parse_cell(text) is None


# ---------------------------------------------------------------- sheet builders
def make_flat(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "BSSE-4A"
    ws.append(["Some title row that is not the header"])
    ws.append(["Day", "Time", "Course", "Teacher", "Room"])
    ws.append(["Monday", "08:30 - 10:00", "SE201 Software Design", "Dr. Ali", "R-2"])
    ws.append(["", "10:00 - 11:30", "CS101 Programming Fundamentals", "Ms. Sara", "R-1"])  # blank day = Monday
    ws.append(["Tuesday", "8:30-10:00", "SE201 Software Design", "Dr. Ali", "R-2"])
    ws.append(["Tuesday", "bad time", "SE201 Software Design", "Dr. Ali", "R-2"])  # unreadable time
    ws.append(["Funday", "08:30 - 10:00", "SE201 Software Design", "Dr. Ali", "R-2"])  # unreadable day
    ws2 = wb.create_sheet("Instructions")
    ws2.append(["Read me"])
    wb.save(path)


def make_grid(path, transposed=False):
    wb = Workbook()
    ws = wb.active
    ws.title = "BSAI-2B"
    slots = ["08:30 - 10:00", "10:00 - 11:30", "11:30 - 1:00"]
    grid = [
        ["Day"] + slots,
        ["Monday", "AI101 Intro to AI\nDr. Noor\nR-1", "Break", "Math\nMs. Zara\nR-2"],
        ["Tuesday", "", "AI101 Intro to AI\nDr. Noor\nR-1", ""],
    ]
    if transposed:
        grid = [list(col) for col in zip(*grid)]
    for row in grid:
        ws.append(row)
    wb.save(path)


# ---------------------------------------------------------------- parsing real files
def test_flat_sheet(tmp_path):
    f = tmp_path / "se.xlsx"
    make_flat(f)
    result = ti.parse_workbook(f, "SE")
    assert result.sheets_read == ["BSSE-4A"]
    assert result.sheets_skipped == ["Instructions"]
    assert len(result.rows) == 3
    first, second, third = result.rows
    assert (first.section, first.day, first.start_time, first.end_time) == ("BSSE-4A", "Monday", time(8, 30), time(10, 0))
    assert (first.course_code, first.course_title, first.teacher, first.room) == ("SE201", "Software Design", "Dr. Ali", "R-2")
    assert second.day == "Monday"  # blank day cell inherited
    assert third.day == "Tuesday"
    messages = " | ".join(i.message for i in result.issues)
    assert "time" in messages and "day" in messages.lower()
    assert all(i.level == "error" for i in result.issues)


def test_grid_sheet(tmp_path):
    f = tmp_path / "ai.xlsx"
    make_grid(f)
    rows = ti.parse_workbook(f, "AI").rows
    assert len(rows) == 3  # the "Break" cell and the empty cell are skipped
    mon_first = next(r for r in rows if r.day == "Monday" and r.start_time == time(8, 30))
    assert (mon_first.course_code, mon_first.course_title, mon_first.teacher, mon_first.room) == (
        "AI101", "Intro to AI", "Dr. Noor", "R-1")
    assert mon_first.section == "BSAI-2B"
    last = next(r for r in rows if r.course_title == "Math")
    assert (last.start_time, last.end_time) == (time(11, 30), time(13, 0))


def test_grid_sheet_transposed_gives_same_result(tmp_path):
    a, b = tmp_path / "a.xlsx", tmp_path / "b.xlsx"
    make_grid(a)
    make_grid(b, transposed=True)
    key = lambda r: (r.day, r.start_time, r.course_title, r.teacher, r.room)  # noqa: E731
    assert sorted(map(key, ti.parse_workbook(a, "CS").rows)) == sorted(map(key, ti.parse_workbook(b, "CS").rows))


def test_merged_day_cells_in_grid(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "S1"
    ws.append(["Day", "Slot", "08:30 - 10:00", "10:00 - 11:30"])
    ws.append(["Monday", "A", "Maths\nDr. X\nR-1", ""])
    ws.append(["", "B", "", "Physics\nDr. Y\nR-2"])
    ws.merge_cells("A2:A3")
    f = tmp_path / "m.xlsx"
    wb.save(f)
    rows = ti.parse_workbook(f, "IT").rows
    assert {(r.day, r.course_title) for r in rows} == {("Monday", "Maths"), ("Monday", "Physics")}


def test_csv_flat(tmp_path):
    f = tmp_path / "BSCS-3A.csv"
    f.write_text("Day,Start,End,Course Code,Course,Teacher,Room\n"
                 "Mon,09:00,10:30,CS301,Algorithms,Dr. Q,R-1\n", encoding="utf-8")
    rows = ti.parse_workbook(f, "CS").rows
    assert len(rows) == 1
    r = rows[0]
    assert (r.section, r.day, r.start_time, r.end_time, r.course_code, r.course_title) == (
        "BSCS-3A", "Monday", time(9, 0), time(10, 30), "CS301", "Algorithms")


def test_unrecognised_sheet_is_reported_not_crashed(tmp_path):
    wb = Workbook()
    wb.active.title = "Weird"
    wb.active.append(["hello", "world"])
    wb.active.append(["foo", "bar"])
    f = tmp_path / "w.xlsx"
    wb.save(f)
    result = ti.parse_workbook(f, "SE")
    assert result.rows == []
    assert any("not recognised" in i.message for i in result.issues)
    assert result.sheets_skipped == ["Weird"]


def test_unknown_department_and_file_type(tmp_path):
    f = tmp_path / "x.csv"
    f.write_text("a,b\n")
    with pytest.raises(ValueError):
        ti.parse_workbook(f, "XX")
    g = tmp_path / "x.txt"
    g.write_text("a")
    with pytest.raises(ValueError):
        ti.parse_workbook(g, "SE")


# ---------------------------------------------------------------- validation
def row(**kw):
    base = dict(department="SE", section="A", day="Monday", start_time=time(9, 0), end_time=time(10, 0),
                course_title="Maths", course_code="M1", source_sheet="A", source_row=2)
    base.update(kw)
    return ti.TimetableRow(**base)


def test_validate_drops_bad_times_and_exact_duplicates():
    rows = [row(), row(), row(end_time=time(9, 0)), row(start_time=time(11, 0), end_time=time(10, 0)),
            row(day="Tuesday")]
    good, issues = ti.validate_rows(rows)
    assert len(good) == 2
    assert sorted(i.level for i in issues) == ["error", "error", "warning"]


# ---------------------------------------------------------------- loading into the database
def test_load_rows_creates_everything_then_is_idempotent(db):
    f_rows = [row(teacher="Dr. Ali", room="R-9"), row(day="Tuesday", teacher="Dr. Ali", room="R-9", source_row=3),
              row(start_time=time(11, 0), end_time=time(12, 0), teacher="TBA", room="", source_row=4)]
    first = ti.load_rows(db, f_rows, "SE")
    db.commit()
    assert first.entries_created == 3
    assert (first.departments_created, first.sections_created, first.courses_created, first.teachers_created) == (1, 1, 1, 1)
    assert first.rooms_created == ["R-9"]
    room = db.execute(select(Room).where(Room.name == "R-9")).scalar_one()
    assert room.capacity == ti.DEFAULT_NEW_ROOM_CAPACITY

    entries = db.execute(select(TimetableEntry).order_by(TimetableEntry.id)).scalars().all()
    assert entries[2].teacher_id is None and entries[2].room_id is None  # TBA / blank -> no teacher / room

    second = ti.load_rows(db, f_rows, "SE")
    db.commit()
    assert second.entries_created == 0
    assert second.entries_skipped_existing == 3
    assert db.execute(select(TimetableEntry)).scalars().all().__len__() == 3


def test_load_rows_reuses_existing_course_and_room(db, seed):
    r = row(course_code="CS101", course_title="Programming Fundamentals", room="R-2", section="BSSE-1A")
    summary = ti.load_rows(db, [r], "SE")
    db.commit()
    assert summary.courses_created == 0
    assert summary.rooms_created == []
    assert summary.departments_created == 0
    assert summary.sections_created == 1
    entry = db.execute(select(TimetableEntry)).scalar_one()
    assert entry.course_id == seed.c1.id and entry.room_id == seed.mid.id
